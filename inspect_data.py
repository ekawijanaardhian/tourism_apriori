import sys
import os

print("Starting inspection...", flush=True)
excel_path = "All - From Intention to Decision - 319 Respondent.xlsx"
print(f"Checking file exists: {os.path.exists(excel_path)}", flush=True)

import openpyxl
print("Loaded openpyxl module...", flush=True)

wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
print(f"Sheet names: {wb.sheetnames}", flush=True)

sheet = wb.active
headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
print(f"Number of columns: {len(headers)}", flush=True)
for i, h in enumerate(headers):
    print(f"{i}: {repr(h)}", flush=True)

row_count = 0
for _ in sheet.iter_rows():
    row_count += 1
print(f"Total rows in sheet (including header): {row_count}", flush=True)
